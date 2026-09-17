"""research_parse.py — extract text from uploaded research files for the brand brief.

Follows the skill's input-parsing rules: PDF via pdfplumber, DOCX via python-docx,
PPTX via python-pptx (text + speaker notes), XLSX via openpyxl, CSV/TXT/MD as text.
Returns a list of {filename, kind, text} blobs, each capped so the prompt stays bounded.
"""
from __future__ import annotations

import csv
import os

PER_FILE_CHARS = 8000  # cap per file so the enrichment prompt stays bounded


def _cap(s: str) -> str:
    s = (s or "").strip()
    return s[:PER_FILE_CHARS]


def _pdf(path: str) -> str:
    import pdfplumber
    out = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages[:30]:
            out.append(page.extract_text() or "")
    return "\n".join(out)


def _docx(path: str) -> str:
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"[Slide {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                parts.append(f"(notes) {note}")
    return "\n".join(parts)


def _xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets[:4]:
        parts.append(f"[Sheet {ws.title}]")
        for r in ws.iter_rows(values_only=True):
            cells = [str(c) for c in r if c is not None]
            if cells:
                parts.append(" | ".join(cells))
            if len(parts) > 200:
                break
    return "\n".join(parts)


def _csv(path: str) -> str:
    with open(path, newline="", encoding="utf-8", errors="ignore") as f:
        rows = list(csv.reader(f))
    return "\n".join(" | ".join(r) for r in rows[:200])


def _text(path: str) -> str:
    return open(path, encoding="utf-8", errors="ignore").read()


_DISPATCH = {
    ".pdf": ("pdf", _pdf), ".docx": ("brand/qual doc", _docx), ".doc": ("brand/qual doc", _docx),
    ".pptx": ("competitor deck", _pptx), ".xlsx": ("data", _xlsx), ".xls": ("data", _xlsx),
    ".csv": ("data", _csv), ".txt": ("qual notes", _text), ".md": ("qual notes", _text),
}


def table_rows(path: str, limit: int = 500) -> dict:
    """A spreadsheet as ROWS, not as flattened text.

    `_xlsx` and `_csv` above join every row with " | " because their consumer is a model reading prose.
    A geography-and-budget sheet is not prose: the columns carry the meaning, and re-splitting the
    flattened string would be a second parser guessing where a cell containing " | " ended.

    Returns `{"kind": "table"|"text", "rows": [[...]], "sheet": str, "text": str, "why": str}`.
    `kind` is "text" for anything that is not a spreadsheet — a deck or a PDF has no rows to map, and
    saying so is better than returning one row per paragraph and calling it a table.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in r]
            if any(cells):
                rows.append(cells)
            if len(rows) >= limit:
                break
        return {"kind": "table", "rows": rows, "sheet": ws.title, "text": "",
                "why": "" if rows else "The first sheet has no rows with anything in them."}
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8", errors="ignore") as f:
            rows = [[(c or "").strip() for c in r] for r in csv.reader(f)]
        rows = [r for r in rows if any(r)][:limit]
        return {"kind": "table", "rows": rows, "sheet": "", "text": "",
                "why": "" if rows else "That CSV has no rows with anything in them."}
    blob = parse_file(path)
    return {"kind": "text", "rows": [], "sheet": "", "text": blob.get("text", ""),
            "why": (f"{blob.get('kind', 'that file')} carries prose, not rows — there are no columns "
                    f"to map to cells. The text is handed over for reading instead.")}


def parse_file(path: str) -> dict:
    ext = os.path.splitext(path)[1].lower()
    kind, fn = _DISPATCH.get(ext, ("unknown", None))
    if fn is None:
        return {"filename": os.path.basename(path), "kind": "unsupported", "text": ""}
    try:
        text = _cap(fn(path))
    except Exception as e:
        text = f"[could not parse: {e}]"
    return {"filename": os.path.basename(path), "kind": kind, "text": text}


def parse_paths(paths: list[str]) -> list[dict]:
    return [parse_file(p) for p in paths if p]


# ==========================================================================================
# Map phase (17 Sep round) — real aggregation and structure instead of flatten-then-truncate.
#
# Found live, against a real user test: `parse_file`/`_cap` above flattens a whole workbook or
# deck to text and hard-truncates at PER_FILE_CHARS, which on a real 12-month x 4-state x
# 10-company panel discarded 95%+ of the file — the model saw a partial, non-representative
# slice of ONE sheet and never reached the others at all. This section replaces that path for
# the brief specifically: spreadsheets get REAL pandas aggregation (bounded by the number of
# distinct categories, not by row count, so a 483-row sheet and a 4,830-row sheet compress to
# roughly the same size); decks/docs get their NATIVE structure read first (title vs. body,
# which the file already encodes) rather than guessed at from flattened prose.
#
# Deliberately not docling/torch: checked docling's own PyPI metadata before reaching for it —
# the default install pulls in torch + a downloaded layout-detection model (multi-GB) for PDF
# layout recovery. PPTX/DOCX don't need that at all (their own lightweight `format-pptx`/
# `format-docx` extras are just python-pptx/python-docx under the hood — the same libraries
# already used below). PDF gets a font-size/position heuristic instead (see `_pdf_structured`)
# — right-sized for a PDF that's actually an exported deck (real, selectable text, a large
# top-of-page block that was a PowerPoint title), which is the real-world case named for this:
# presentations shared as PDF for file-size/compatibility reasons, not scanned paper. A page
# with no extractable text layer is flagged `"scanned": true` rather than silently returning
# nothing — OCR is a genuinely different, heavier problem, not built pre-emptively for a case
# that hasn't come up yet.
# ==========================================================================================

import re as _re

_README_NAME = _re.compile(r"read.?me|instructions?|about|notes?|guide|legend|cover", _re.I)
MAX_DISTINCT = 50       # cap on how many distinct values of one categorical column get listed
MAX_GROUP_DIMS = 2      # how many categorical columns get cross-tabulated together
MAX_AGG_ROWS = 400      # backstop on the aggregate table size for a pathological sheet


def _reheader(raw):
    """Real workbooks (both files this was tested against, in fact) commonly put a merged title
    in row 1 and the real column headers a row or two below it — reading row 0 as the header (the
    default) turns every real column into "Unnamed: N" and the actual headers into a data row.
    Pick whichever of the first 8 rows has the most non-null cells as the header — a title row
    (one cell filled) and a stray note row lose to the real header row (every column filled) under
    that rule, on both real layouts this was checked against."""
    if raw.empty:
        return raw
    window = raw.head(8)
    counts = window.notna().sum(axis=1)
    header_idx = int(counts.idxmax())
    if counts.loc[header_idx] < 2:
        return raw  # nothing looks like a real header; leave the sheet as-is
    header = raw.iloc[header_idx].tolist()
    body = raw.iloc[header_idx + 1:].reset_index(drop=True)
    names = [str(h).strip() if h is not None and str(h).strip() else f"col_{i}"
             for i, h in enumerate(header)]
    # Deduplicate — a validation/summary sub-table sitting beside the main one (again, the real
    # shape both test files use) commonly repeats a header like "Month" a second time; pandas
    # can't assign that back as columns without every name being unique.
    seen: dict[str, int] = {}
    deduped = []
    for n in names:
        seen[n] = seen.get(n, 0) + 1
        deduped.append(n if seen[n] == 1 else f"{n} ({seen[n]})")
    body.columns = deduped
    return body


def _sheet_is_readme(name: str, df) -> bool:
    if _README_NAME.search(name or ""):
        return True
    # A sheet that's almost entirely one text column with few real data rows reads as a cover/
    # notes page, not data — same shape as the README tabs in the files this was built against.
    return len(df) < 3 and df.shape[1] <= 2


def _classify_columns(df):
    """Split a dataframe's columns into {time, dims, metrics, ignored} by shape, not by name —
    this has to work on a spreadsheet whose columns we've never seen, not just the ones tested
    against. A time column is whatever parses cleanly as dates; a metric is numeric; a dimension
    is text/categorical with low cardinality relative to the row count; anything else (a free-text
    or near-unique id column) is ignored rather than guessed at."""
    import pandas as pd

    time_col, dims, metrics, ignored = None, [], [], []
    n = max(1, len(df))
    for col in df.columns:
        s = df[col]
        name = str(col)
        if time_col is None:
            parsed = pd.to_datetime(s, errors="coerce")
            if parsed.notna().mean() > 0.9:
                time_col = name
                continue
        numeric = pd.to_numeric(s, errors="coerce")
        if numeric.notna().mean() > 0.9:
            metrics.append(name)
            continue
        nunique = s.nunique(dropna=True)
        # >=2, not >=1: a column with a single distinct value (a constant validation column like
        # a "Check (=100%?)" tab, or a sparse column that's mostly-blank in a wide re-headered
        # sheet with a validation table's own short column sitting beside the main one) carries
        # zero grouping information — a groupby on it is a no-op, so it's noise, not a dimension.
        if 2 <= nunique <= min(MAX_DISTINCT, max(10, int(n * 0.5))):
            dims.append(name)
        else:
            ignored.append(name)
    return time_col, dims, metrics, ignored


def analyze_spreadsheet(path: str, max_sheets: int = 6) -> dict:
    """Real aggregation over an uploaded xlsx/csv — mean/min/max per dimension combination, a
    trend read where a time column exists, correlation between metric pairs — computed with
    pandas/scipy, not eyeballed by a model reading flattened rows. This is also what closes the
    accuracy gap the flatten approach had: a model asked to average hundreds of raw text rows
    itself produced a visibly-off range in the live test this was built from; a real `groupby().
    mean()` doesn't.

    Returns {"filename", "kind": "data", "sheets": [{"name", "rows", skipped_why?, "dims": [...],
    "metrics": [...], "distinct": {dim: [values]}, "aggregates": [...], "trend": [...],
    "correlations": [...]}], "why": str}
    """
    import pandas as pd

    filename = os.path.basename(path)
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".csv":
            raw = pd.read_csv(path, header=None)
            sheets = {"Sheet1": _reheader(raw)}
        else:
            xls = pd.ExcelFile(path, engine="openpyxl")
            names = xls.sheet_names[: max_sheets + 3]  # a little headroom for README tabs
            sheets = {n: _reheader(xls.parse(n, header=None)) for n in names}
    except Exception as e:
        return {"filename": filename, "kind": "data", "sheets": [],
                "why": f"could not read as a spreadsheet: {type(e).__name__}: {e}"}

    out_sheets = []
    used = 0
    for name, df in sheets.items():
        if used >= max_sheets:
            out_sheets.append({"name": name, "rows": len(df), "skipped_why":
                               "sheet cap reached — see max_sheets"})
            continue
        df = df.dropna(how="all")
        if df.empty:
            continue
        if _sheet_is_readme(name, df):
            out_sheets.append({"name": name, "rows": len(df), "skipped_why":
                               "reads as a cover/notes/readme tab, not data"})
            continue
        try:
            entry = _analyze_one_sheet(name, df)
        except Exception as e:
            # A second, validation-style table sitting beside the main one on the same sheet (the
            # shape both real test files use — a SUMIFS check block to the right of the data) can
            # produce a column that's genuinely mixed-type once pandas reads the whole sheet width.
            # One sheet failing this way shouldn't lose the rest of the file's real data.
            entry = {"name": name, "rows": len(df), "skipped_why":
                     f"could not analyze this sheet: {type(e).__name__}: {e}"}
        used += 1
        out_sheets.append(entry)

    return {"filename": filename, "kind": "data", "sheets": out_sheets, "why": ""}


def _pick_group_dims(df, dims: list[str], k: int = MAX_GROUP_DIMS) -> list[str]:
    """Greedily pick up to `k` dimension columns to cross-tabulate, lowest cardinality first, but
    skipping anything that's really a relabeling of a dimension already picked — a real live-tested
    case: "State" and "State Code" both have 4 distinct values and both look equally attractive by
    cardinality alone, but they carry the SAME information twice ("Andhra Pradesh" / "AP"), so
    picking both wastes half the grouping budget on a no-op instead of on "Company," the dimension
    that actually varies independently and is the one worth cross-tabulating against."""
    import pandas as pd

    ranked = sorted(dims, key=lambda d: df[d].nunique())
    picked: list[str] = []
    for d in ranked:
        if len(picked) >= k:
            break
        if picked:
            try:
                avg_within = df.groupby(picked, dropna=True)[d].nunique().mean()
            except Exception:
                avg_within = 2  # if the check itself fails, don't let it block a real pick
            if pd.notna(avg_within) and avg_within <= 1.05:
                continue  # near-bijection with what's already picked — same dimension, skip it
        picked.append(d)
    return picked


def _analyze_one_sheet(name: str, df) -> dict:
    import pandas as pd

    time_col, dims, metrics, _ignored = _classify_columns(df)
    if time_col:
        # Force a homogeneous datetime dtype now — a mixed-type column (e.g. a second table's own
        # "Month" header sitting a few columns over on the same sheet) otherwise crashes sort_values
        # later with real dates next to stray strings in the same column.
        df = df.copy()
        df[time_col] = pd.to_datetime(df[time_col], errors="coerce")

    entry = {"name": name, "rows": len(df), "dims": dims, "metrics": metrics}
    if time_col:
        entry["time_column"] = time_col

    # Full distinct-value list per dimension — the fix for the specific live-tested gap where a
    # real competitor set got silently narrowed by wherever row-flattening happened to stop. This
    # is exact and complete regardless of file size.
    entry["distinct"] = {
        d: sorted(str(v) for v in df[d].dropna().unique())[:MAX_DISTINCT] for d in dims
    }

    group_dims = _pick_group_dims(df, dims)
    entry["cross_tabulated_on"] = group_dims
    if dims and not group_dims:
        entry["also_present"] = [d for d in dims if d not in group_dims]

    aggregates = []
    if group_dims and metrics:
        for key, sub in df.groupby(group_dims, dropna=True):
            key_t = key if isinstance(key, tuple) else (key,)
            row = dict(zip(group_dims, (str(k) for k in key_t)))
            for m in metrics:
                vals = pd.to_numeric(sub[m], errors="coerce").dropna()
                if len(vals):
                    row[m] = {"mean": round(float(vals.mean()), 4), "min": round(float(vals.min()), 4),
                              "max": round(float(vals.max()), 4), "n": int(len(vals))}
            aggregates.append(row)
            if len(aggregates) >= MAX_AGG_ROWS:
                break
    entry["aggregates"] = aggregates

    trend = []
    if time_col and group_dims and metrics:
        df = df.sort_values(time_col)
        for key, sub in df.groupby(group_dims, dropna=True):
            key_t = key if isinstance(key, tuple) else (key,)
            for m in metrics:
                vals = pd.to_numeric(sub[m], errors="coerce").dropna()
                if len(vals) < 3:
                    continue
                x = list(range(len(vals)))
                try:
                    from scipy import stats as _stats
                    slope = _stats.linregress(x, vals.tolist()).slope
                except Exception:
                    slope = (vals.iloc[-1] - vals.iloc[0]) / max(1, len(vals) - 1)
                span = float(vals.max() - vals.min()) or 1.0
                rel = slope * len(vals) / span
                direction = "rising" if rel > 0.15 else "falling" if rel < -0.15 else "flat"
                trend.append({**dict(zip(group_dims, (str(k) for k in key_t))),
                              "metric": m, "direction": direction,
                              "from": round(float(vals.iloc[0]), 4), "to": round(float(vals.iloc[-1]), 4)})
                if len(trend) >= MAX_AGG_ROWS:
                    break
    entry["trend"] = trend

    correlations = []
    if len(metrics) >= 2:
        try:
            from scipy import stats as _stats
            for i in range(len(metrics)):
                for j in range(i + 1, len(metrics)):
                    a = pd.to_numeric(df[metrics[i]], errors="coerce")
                    b = pd.to_numeric(df[metrics[j]], errors="coerce")
                    both = pd.concat([a, b], axis=1).dropna()
                    if len(both) >= 5:
                        r, p = _stats.pearsonr(both.iloc[:, 0], both.iloc[:, 1])
                        if abs(r) >= 0.3:
                            correlations.append({"a": metrics[i], "b": metrics[j],
                                                 "r": round(float(r), 3), "p": round(float(p), 4),
                                                 "n": int(len(both))})
        except Exception:
            pass
    entry["correlations"] = correlations
    return entry


def _pptx_structured(path: str) -> list[dict]:
    """[{slide, title, body, tables, notes}] — reads the placeholder type PowerPoint already
    assigns (title vs. body) instead of flattening every text frame equally. This is the same
    information Docling's own `format-pptx` backend would read (it's built on python-pptx too),
    without the layout-model weight that format doesn't need.

    Live-tested finding: a deck built without real placeholders at all (every shape reports
    `is_placeholder=False`, e.g. built programmatically rather than from a PowerPoint template —
    the real qualitative-research test file this was checked against is exactly this shape) has
    nothing for the placeholder/name checks to find. Font size turned out NOT to be a reliable
    fallback either — checked directly against that file, and a subtitle line came back sized
    larger than the actual headline. Position is: the topmost shape on the slide, checked directly
    against the same file, is reliably the title even with zero placeholder metadata — same
    "topmost = title" principle the PDF heuristic below already uses."""
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER_TYPE as PT

    _TITLE_TYPES = {PT.TITLE, PT.CENTER_TITLE}
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        title, body_parts, tables = "", [], []
        text_shapes = []  # (top, text) for the position-based fallback
        placeholder_title = ""
        for shape in slide.shapes:
            if shape.has_table:
                tbl = shape.table
                tables.append([[c.text.strip() for c in row.cells] for row in tbl.rows])
                continue
            if not (shape.has_text_frame and shape.text_frame.text.strip()):
                continue
            text = shape.text_frame.text.strip()
            is_title = False
            try:
                ph = shape.placeholder_format
                is_title = ph is not None and ph.type in _TITLE_TYPES
            except Exception:
                pass
            if not is_title and "title" in (shape.name or "").lower():
                is_title = True
            if is_title and not placeholder_title:
                placeholder_title = text
            else:
                body_parts.append(text)
            if shape.top is not None:
                text_shapes.append((shape.top, text))
        if placeholder_title:
            title = placeholder_title
        elif text_shapes:
            # No placeholder/name signal at all — fall back to whichever text shape sits highest
            # on the slide, and pull it back out of body_parts so it isn't duplicated.
            text_shapes.sort(key=lambda t: t[0])
            title = text_shapes[0][1]
            if title in body_parts:
                body_parts.remove(title)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        out.append({"slide": i, "title": title, "body": "\n".join(body_parts),
                    "tables": tables, "notes": notes})
    return out


def _docx_structured(path: str) -> dict:
    """{"sections": [{"heading_level", "heading", "body": [...]}], "tables": [[...]]} — reads
    the paragraph style name (Heading 1/2/…) Word already assigns instead of treating every
    paragraph as equally flat prose."""
    import docx

    d = docx.Document(path)
    sections, tables = [], []
    current = {"heading_level": 0, "heading": "", "body": []}
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "") if p.style else ""
        low = style.lower()
        if low.startswith("heading") or low in ("title", "subtitle"):
            if current["heading"] or current["body"]:
                sections.append(current)
            m = _re.search(r"(\d+)", style)
            current = {"heading_level": int(m.group(1)) if m else 0, "heading": text, "body": []}
        else:
            current["body"].append(text)
    if current["heading"] or current["body"]:
        sections.append(current)
    for t in d.tables:
        tables.append([[c.text.strip() for c in row.cells] for row in t.rows])
    return {"sections": sections, "tables": tables}


def _pdf_structured(path: str, max_pages: int = 60) -> list[dict]:
    """[{page, title, body, tables, scanned}] — a font-size/position heuristic standing in for
    PDF's total lack of native structure: the largest-font text block in the top ~30% of a page
    is treated as that page's title, matching how a PowerPoint title looks once exported (large,
    top-positioned) — right-sized for a real deck exported to PDF, not a scanned document, which
    has no text layer at all and gets flagged `scanned: true` rather than guessed at."""
    import pdfplumber

    out = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages[:max_pages], 1):
            chars = page.chars or []
            text = (page.extract_text() or "").strip()
            if not text and not chars:
                out.append({"page": i, "title": "", "body": "", "tables": [], "scanned": True})
                continue
            title, body = "", text
            if chars:
                lines: dict[int, list] = {}
                for c in chars:
                    lines.setdefault(round(c["top"]), []).append(c)
                line_items = []
                for top, cs in lines.items():
                    t = "".join(c["text"] for c in sorted(cs, key=lambda c: c["x0"])).strip()
                    if t:
                        line_items.append((top, max(c.get("size", 0) for c in cs), t))
                line_items.sort(key=lambda x: x[0])
                if line_items:
                    page_h = getattr(page, "height", 0) or max(t for t, _, _ in line_items) or 1
                    candidates = [li for li in line_items if li[0] < page_h * 0.3] or line_items[:3]
                    best = max(candidates, key=lambda x: x[1])
                    title = best[2]
                    body = "\n".join(t for (_, _, t) in line_items if t != title)
            tables = []
            try:
                tables = page.extract_tables() or []
            except Exception:
                pass
            out.append({"page": i, "title": title, "body": body, "tables": tables, "scanned": False})
    return out


_DOC_DISPATCH = {".pptx": ("pptx", _pptx_structured), ".docx": ("docx", _docx_structured),
                 ".pdf": ("pdf", _pdf_structured)}


def analyze_document(path: str) -> dict:
    """Structured extraction for a deck/doc/PDF — {"filename", "kind": "deck", "format", "content"}.
    `content` is a pptx-style list of section dicts for every format (docx's own {"sections",
    "tables"} shape gets flattened to the same {slide/section-index, title, body, tables} list so
    downstream code doesn't need a format branch)."""
    filename = os.path.basename(path)
    ext = os.path.splitext(path)[1].lower()
    fmt_fn = _DOC_DISPATCH.get(ext)
    if not fmt_fn:
        return {"filename": filename, "kind": "unsupported", "format": ext, "content": []}
    fmt, fn = fmt_fn
    try:
        raw = fn(path)
    except Exception as e:
        return {"filename": filename, "kind": "deck", "format": fmt, "content": [],
                "why": f"could not parse: {type(e).__name__}: {e}"}
    if fmt == "docx":
        content = [{"idx": i + 1, "title": s["heading"], "body": "\n".join(s["body"]), "tables": []}
                   for i, s in enumerate(raw["sections"])]
        if raw["tables"]:
            content.append({"idx": len(content) + 1, "title": "(tables)", "body": "",
                            "tables": raw["tables"]})
    else:
        content = [{"idx": s.get("slide") or s.get("page"), "title": s["title"], "body": s["body"],
                   "tables": s["tables"], **({"scanned": s["scanned"]} if "scanned" in s else {})}
                  for s in raw]
    return {"filename": filename, "kind": "deck", "format": fmt, "content": content}


def analyze_upload(path: str) -> dict:
    """Top-level Map-phase dispatcher for one uploaded file — real aggregation for spreadsheets,
    structured extraction for decks/docs/PDF. Unknown formats fall back to the old flat `parse_file`
    so nothing that worked before regresses."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xls", ".csv"):
        return analyze_spreadsheet(path)
    if ext in _DOC_DISPATCH:
        return analyze_document(path)
    blob = parse_file(path)
    return {"filename": blob["filename"], "kind": "text", "format": ext, "content": blob["text"]}


_CHART_IMAGE_MEDIA = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
                      "gif": "image/gif", "webp": "image/webp"}


def chart_image_from_upload(path: str) -> dict | None:
    """For a NeedScope wheel / CB-CA reference chart upload: return {"media_type", "data_b64"} if the
    file can be read as an image the model can see — either because it already is one, or because it's
    a PDF (rasterize page 1 — pypdfium2 is already a dependency via pdfplumber, so this needed nothing
    new). Returns None for anything else: a PPTX/DOCX chart built from native shapes has no single image
    to extract without a real rendering engine (LibreOffice) this deployment's instance size can't
    safely carry yet — see BRAND_GROUNDING_TESTING_LOG.md's entry on why that was deliberately deferred
    rather than risking the whole app's memory budget. Callers should fold a None result into the
    regular research pipeline instead of silently dropping it — the chart's raw content still reaches
    the brief that way, just not as the wheel/CB-CA reading specifically.
    """
    import base64
    ext = os.path.splitext(path)[1].lower().lstrip(".")
    if ext in _CHART_IMAGE_MEDIA:
        with open(path, "rb") as f:
            return {"media_type": _CHART_IMAGE_MEDIA[ext], "data_b64": base64.b64encode(f.read()).decode("ascii")}
    if ext == "pdf":
        try:
            import io
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(path)
            if len(pdf) == 0:
                return None
            pil_img = pdf[0].render(scale=2.0).to_pil().convert("RGB")
            buf = io.BytesIO()
            pil_img.save(buf, format="PNG")
            return {"media_type": "image/png", "data_b64": base64.b64encode(buf.getvalue()).decode("ascii")}
        except Exception:
            return None
    return None


# ---------------------------------------------------------------------------------------------------
# Map, second half: bounded per-file LLM extraction for decks/docs/PDF.
#
# The spreadsheet half (analyze_spreadsheet) gets real numbers from real code — groupby/mean/trend/
# correlation, computed, not guessed. A slide deck has no such structure to compute over; the only way
# to get "what does this document actually say" out of free text is to read it. But reading it must not
# mean re-inflating the old truncate-and-dump bug at a different layer — the user's correction was
# explicit: "the idea is to find patterns across data using analytical techniques and not necessarily
# summarise everything that's thrown at you." So this call is bounded on BOTH ends: bounded input (a
# title list up front so the model sees the whole document's shape before any body text, then capped
# body/table text per section), and bounded output (a short, fixed-shape findings list regardless of
# whether the source is a 3-slide deck or a 60-slide one) — a 60-slide deck does not get a 10x longer
# answer, it gets a more selective one. That selectivity is the point: the prompt explicitly asks the
# model to skip agenda/methodology/thank-you slides and to only keep what a senior marketer would
# actually cite, so most of a deck legitimately produces nothing and that is reported, not padded out.
# ---------------------------------------------------------------------------------------------------

_SECTION_CHARS = 700       # per-section body cap — generous enough for a real paragraph, not a dump
_SECTION_TABLE_ROWS = 12   # per-table row cap when a section carries a table
_MAX_PROMPT_CHARS = 24000  # overall content budget for the prompt body, independent of section count


def _sections_for_prompt(content: list[dict]) -> tuple[str, list[int]]:
    """Render section bodies/tables for the extraction prompt, capped per-section and overall.
    Returns (rendered_text, idxs_dropped_for_budget) so the prompt can be honest that some sections
    were not shown at all, rather than silently omitting them."""
    parts, used, dropped = [], 0, []
    for s in content:
        idx = s.get("idx")
        title = s.get("title") or "(untitled)"
        body = (s.get("body") or "").strip()
        if len(body) > _SECTION_CHARS:
            body = body[:_SECTION_CHARS] + " …[truncated]"
        table_txt = ""
        tables = s.get("tables") or []
        if tables:
            t = tables[0]
            rows = t[:_SECTION_TABLE_ROWS]
            table_txt = "\nTable:\n" + "\n".join(" | ".join(str(c) for c in row) for row in rows)
            if len(tables) > 1 or len(t) > _SECTION_TABLE_ROWS:
                table_txt += "\n[additional table rows/tables omitted for length]"
        notes = (s.get("notes") or "").strip()
        note_txt = f"\nSpeaker notes: {notes[:300]}" if notes else ""
        flag = " [scanned page — no extractable text]" if s.get("scanned") else ""
        block = f"--- [{idx}] {title}{flag} ---\n{body}{table_txt}{note_txt}"
        if used + len(block) > _MAX_PROMPT_CHARS:
            dropped.append(idx)
            continue
        used += len(block)
        parts.append(block)
    return "\n\n".join(parts), dropped


_EXTRACT_PROMPT = """You are extracting decision-relevant findings from ONE uploaded research document, \
as part of building the evidence base for a marketing brief.{focus_line}

Below are ALL section titles in this document in order, so you can see its whole shape before reading \
any body text:
{titles}

Full section content follows (some very long sections are capped; tables and speaker notes are included \
where present):
{body}
{dropped_line}
Extract ONLY what a senior marketer would actually cite in a brief: specific numbers, named findings, \
named frameworks/models, quotes worth quoting verbatim, and explicit recommendations. Skip agenda \
slides, methodology-only slides, title/thank-you slides, and section headers with nothing under them — \
list those section numbers under "skipped_sections" instead of forcing something out of them. Only \
extract what is actually written above — never infer, generalize, or fill a gap with outside knowledge. \
Return AT MOST 25 findings — if more qualify, keep only the 25 most decision-relevant and fold the rest \
into coverage_note rather than listing every one; keep each claim to one sentence.

Return ONLY this JSON object:
{{
  "findings": [
    {{"claim": "one sentence, specific, citable", "kind": "finding|quote|framework|recommendation|number", \
"source_idx": <int from the section markers above>, "source_label": "<its title>", "verbatim": true|false}}
  ],
  "skipped_sections": [<idx>, ...],
  "coverage_note": "one honest sentence: what this document is mostly about, and any major gap you noticed"
}}"""


def summarize_document(analyzed: dict, *, focus: str = "") -> dict:
    """Bounded per-file LLM extraction — the second half of Map. Input is the dict `analyze_document()`
    (or the `.content` shape of `analyze_upload()` for a deck/doc/PDF) produced; output is always a
    short, fixed-shape findings list — this is deliberately NOT a summary of the whole file, it is a
    selection from it, tagged back to the section/slide/page it came from so a later citation can point
    at something real. `focus` is optional context (the brief's brand/category/ask) that helps the model
    prioritize what to keep; it is a hint, not a filter — the extraction still only pulls what is
    actually present in the text, per the "never invent" rule this whole pipeline is held to."""
    filename = analyzed.get("filename", "")
    content = analyzed.get("content") or []
    if not content:
        return {"filename": filename, "findings": [], "skipped_sections": [],
                "coverage_note": analyzed.get("why") or "no extractable content"}

    titles = "\n".join(f"[{s.get('idx')}] {s.get('title') or '(untitled)'}" for s in content)
    body, dropped = _sections_for_prompt(content)
    focus_line = (f"\nThis research is being read to inform a marketing brief for: {focus}\n"
                  if focus else "")
    dropped_line = (f"\n[Sections {dropped} were not shown above — the document exceeded this pass's "
                     f"content budget; treat them as unread, not as containing nothing.]\n"
                     if dropped else "")
    prompt = _EXTRACT_PROMPT.format(focus_line=focus_line, titles=titles, body=body,
                                    dropped_line=dropped_line)

    from jsonout import ask_json
    # 25 findings * (~1 sentence claim + label/kind/idx fields) fits well inside 3200 tokens even with
    # JSON formatting overhead — sized against the real cut-off seen with a 16-slide deck at 1800.
    data, err = ask_json(prompt, max_tokens=3200)
    if data is None:
        return {"filename": filename, "findings": [], "skipped_sections": [],
                "coverage_note": f"extraction failed: {err}"}
    data["filename"] = filename
    data.setdefault("findings", [])
    data.setdefault("skipped_sections", [])
    data.setdefault("coverage_note", "")
    if dropped:
        data["skipped_sections"] = sorted(set(data["skipped_sections"]) | set(dropped))
    return data


# ---------------------------------------------------------------------------------------------------
# Synthesize: cross-file relevance, corroboration, and honest disagreement.
#
# Map produces one small, bounded result per file — real aggregates for spreadsheets, a short findings
# list for decks/docs/PDF. Synthesize is the step the user asked for by name twice: "not everything is
# given equal weightage, some may not be considered at all, while similar findings can be merged as
# well" — so this is not a second summarization pass, it is a cross-file comparison. Two findings from
# two different files that say the same thing become ONE claim with both sources attached and a
# corroboration count; a claim only one file supports stays single-sourced and says so; a claim that
# directly contradicts another is kept as a flagged disagreement with both sides named, never silently
# resolved in either direction; and anything that turns out irrelevant to this brief's brand/category/
# ask is listed under "considered_not_used" with a one-line reason, rather than dropped without a trace
# — the same "say what's missing, never invent" discipline this pipeline is held to everywhere else,
# applied here to *why something was left out* rather than only to what was kept in.
# ---------------------------------------------------------------------------------------------------

_MAX_HIGHLIGHTS_PER_SHEET = 8


def _spreadsheet_highlights(sheet: dict) -> list[str]:
    """Turn one analyze_spreadsheet() sheet entry into a short list of plain-English bullets — the
    same compaction the deck side gets from summarize_document, so Synthesize sees a uniform shape
    regardless of source type. Picks the extremes (highest/lowest group averages) rather than listing
    every aggregate row, because a synthesis prompt that re-lists 400 groupby rows per sheet across up
    to 20 files is exactly the flat-dump failure mode this whole pipeline replaced."""
    if sheet.get("skipped_why"):
        return []
    name = sheet.get("name", "")
    dims = sheet.get("cross_tabulated_on") or []
    metrics = sheet.get("metrics") or []
    out = []
    aggs = sheet.get("aggregates") or []
    if aggs and dims and metrics:
        m = metrics[0]
        scored = [(row, row[m]["mean"]) for row in aggs if isinstance(row.get(m), dict)]
        scored.sort(key=lambda t: t[1], reverse=True)
        top = scored[:4]
        bottom = [t for t in scored[-3:] if t not in top]
        for row, val in top + bottom:
            label = ", ".join(f"{d}={row.get(d)}" for d in dims)
            out.append(f"[{name}] by {label}: {m} averages {val:g} (n={row[m]['n']})")
    for t in (sheet.get("trend") or [])[:6]:
        label = ", ".join(f"{k}={v}" for k, v in t.items() if k not in ("metric", "direction", "from", "to"))
        out.append(f"[{name}] {label} — {t['metric']} is {t['direction']} "
                   f"(from {t['from']:g} to {t['to']:g})")
    for c in (sheet.get("correlations") or [])[:4]:
        out.append(f"[{name}] {c['a']} and {c['b']} move together: r={c['r']} (n={c['n']})")
    return out[:_MAX_HIGHLIGHTS_PER_SHEET]


def _file_block_for_synthesis(result: dict) -> tuple[str, int]:
    """One file's Map output → (rendered block for the synthesis prompt, item count) — the item count
    lets the caller report, per file, how much of it actually had anything to say."""
    filename = result.get("filename", "?")
    if result.get("kind") == "data":
        lines = []
        for sheet in result.get("sheets") or []:
            lines.extend(_spreadsheet_highlights(sheet))
            why = sheet.get("skipped_why")
            if why:
                lines.append(f"[{sheet.get('name')}] not analyzed — {why}")
        if not lines:
            lines = [result.get("why") or "no analyzable data found in this file"]
        body = "\n".join(f"- {l}" for l in lines)
        return f"=== {filename} (spreadsheet) ===\n{body}", len(lines)
    findings = result.get("findings") or []
    if not findings:
        note = result.get("coverage_note") or "no citable findings extracted"
        return f"=== {filename} ===\n- {note}", 0
    lines = [f"- ({f.get('kind')}) {f.get('claim')} [source: {f.get('source_label')}]" for f in findings]
    cov = result.get("coverage_note", "")
    body = "\n".join(lines) + (f"\n(document overview: {cov})" if cov else "")
    return f"=== {filename} ===\n{body}", len(findings)


_SYNTHESIZE_PROMPT = """You are a senior marketing analyst finding cross-source patterns across several \
research documents, to inform ONE marketing brief.{focus_line}

Below is a short, already-distilled set of findings/highlights from each uploaded document (already \
extracted from much longer originals — treat each line as reliable, not as something to re-verify):

{blocks}

Your job is analysis across these documents, not another summary of each one individually:
1. Where two or more documents say essentially the same thing, MERGE them into one claim and list every \
supporting document.
2. Where documents genuinely disagree or point opposite directions, keep BOTH sides as a flagged \
disagreement — never silently pick a winner.
3. Rank claims by how decision-relevant they are to this specific brief. A claim only one weak or \
tangential source supports should usually NOT become a headline claim.
4. If a document (or part of one) is off-topic, redundant, or too thin to use, say so explicitly under \
"considered_not_used" with a one-line reason — do not just leave it out silently.
5. If the documents do NOT cohere into one clean story, say that plainly in "pattern_note" rather than \
forcing a narrative that isn't really there — competing or inconclusive reads are an acceptable answer.

Return ONLY this JSON object:
{{
  "claims": [
    {{"claim": "one or two sentences, specific and citable", "supporting_sources": ["filename", ...], \
"dissenting_sources": ["filename", ...], "confidence": "high|medium|low", \
"why": "why this matters for the brief, one sentence"}}
  ],
  "considered_not_used": [
    {{"source": "filename", "reason": "one sentence"}}
  ],
  "pattern_note": "one honest paragraph: what coheres across sources, what doesn't, and any real \
disagreement — say plainly if there is no clean single pattern"
}}"""


def synthesize_research(file_results: list[dict], *, focus: str = "") -> dict:
    """Synthesize phase: takes a list of per-file Map outputs (analyze_spreadsheet() or
    summarize_document() results, one per uploaded file) and produces ONE cross-file synthesis —
    merged/corroborated claims, flagged disagreements, and an explicit "considered, not used" list, so
    Reduce (the brief-writing prompt) consumes a point of view instead of N independent file summaries."""
    if not file_results:
        return {"claims": [], "considered_not_used": [], "pattern_note": "no research files were provided"}

    blocks, empties = [], []
    for r in file_results:
        block, n_items = _file_block_for_synthesis(r)
        blocks.append(block)
        if n_items == 0:
            empties.append(r.get("filename", "?"))

    focus_line = f"\nThis brief is for: {focus}\n" if focus else ""
    prompt = _SYNTHESIZE_PROMPT.format(focus_line=focus_line, blocks="\n\n".join(blocks))

    from jsonout import ask_json
    # Was 4000 — proven too small with a live 10-file run (the exact scaling case Phase 1 was built
    # for): the reply was cut off mid-JSON, silently losing the merge/confidence/considered-not-used
    # layer to the unmerged-claims fallback below. Sized up so the ~20-file case this pipeline was
    # designed for has real headroom, not just the file count already tested.
    data, err = ask_json(prompt, max_tokens=7000)
    if data is None:
        # Synthesis failing must not silently mean "no research was used" — fall back to a flat,
        # unmerged claims list built straight from the Map outputs so the brief still has citable
        # material, just without the cross-file merge/dissent layer.
        claims = []
        for r in file_results:
            for f in (r.get("findings") or []):
                claims.append({"claim": f.get("claim", ""), "supporting_sources": [r.get("filename", "?")],
                               "dissenting_sources": [], "confidence": "low",
                               "why": "synthesis step failed — shown unmerged"})
        return {"claims": claims, "considered_not_used": [],
                "pattern_note": f"cross-file synthesis failed ({err}); showing per-file findings unmerged"}

    data.setdefault("claims", [])
    data.setdefault("considered_not_used", [])
    data.setdefault("pattern_note", "")
    if empties:
        already = {c.get("source") for c in data["considered_not_used"]}
        for fn in empties:
            if fn not in already:
                data["considered_not_used"].append({"source": fn, "reason": "no citable findings in Map phase"})
    return data


# ---------------------------------------------------------------------------------------------------
# Reduce, entry point: one call that runs Map then Synthesize over a brief's whole upload set.
#
# Both callers that mine uploaded research (brief_ai.py's AI-drafting step and brandbrief.py's
# doc-export enrichment) used to each call research_parse.parse_paths() and do their own raw-text
# truncation independently — which is how the original bug (8000 chars/file, 16000 combined) ended up
# duplicated in two places. This is the one function both should call instead, so there is exactly one
# ingestion pipeline, not two that can drift out of sync.
# ---------------------------------------------------------------------------------------------------

def _map_one_file(p: str, focus: str) -> dict:
    """One file's whole Map step — structural extraction, then (for decks) the bounded LLM extraction
    call. Split out of ingest_for_brief() so it can run in a worker thread: each call is independent
    (its own local parse, its own fresh anthropic.Anthropic() client inside ask_json — no shared mutable
    state), so N files in parallel is safe, not just faster.

    Wrapped in one blanket try/except: running under ThreadPoolExecutor.map(), an uncaught exception
    here would surface when the results are consumed and abort the WHOLE batch for every other file,
    not just this one — worse under concurrency than the old sequential loop, where at least the files
    before the crash had already been appended. One bad file must degrade to an error entry, never take
    the others down with it.
    """
    try:
        try:
            analyzed = analyze_upload(p)
        except Exception as e:
            return {"filename": os.path.basename(p), "kind": "error", "findings": [],
                    "why": f"{type(e).__name__}: {e}"}
        if analyzed.get("kind") == "deck":
            return summarize_document(analyzed, focus=focus)
        if analyzed.get("kind") == "data":
            return analyzed
        # unsupported/plain-text fallback — analyze_upload's own fallback shape ({kind:"text",
        # content:str}); give Synthesize something to at least name and skip cleanly.
        text = str(analyzed.get("content") or "")[:PER_FILE_CHARS]
        return {"filename": analyzed.get("filename", os.path.basename(p)),
                "findings": ([{"claim": text[:280], "kind": "finding", "source_idx": 1,
                              "source_label": "(plain text)", "verbatim": False}] if text.strip() else []),
                "coverage_note": "" if text.strip() else "no extractable text"}
    except Exception as e:
        return {"filename": os.path.basename(p), "kind": "error", "findings": [],
                "why": f"unexpected failure processing this file: {type(e).__name__}: {e}"}


def ingest_for_brief(paths: list[str], *, focus: str = "") -> dict:
    """Run Map (per-file analysis/extraction) then Synthesize (cross-file merge/dissent) over every
    uploaded file for one brief. Returns {"filenames": [...], "synth": {claims, considered_not_used,
    pattern_note}, "file_results": [...]} — `file_results` is kept for callers that also want the raw
    per-file spreadsheet aggregates (e.g. for a verbatim market-data table), `synth` is what the
    brief-writing prompts should actually read.

    Map runs one worker thread per file (capped) rather than sequentially — a real production 524
    (Cloudflare's proxy timing out waiting for the origin) surfaced this: 5 decks meant 5 sequential
    bounded LLM calls plus Synthesize plus the draft call itself, comfortably past the ~100s a proxy
    will wait. Running the independent per-file calls concurrently turns that from "sum of every file's
    call" into "the slowest single file's call," which is the only lever available here short of making
    the whole route asynchronous (poll-for-result) — a bigger change than this bug needed.
    """
    filenames = [os.path.basename(p) for p in paths]
    if not paths:
        return {"filenames": filenames, "synth": synthesize_research([], focus=focus), "file_results": []}
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=min(8, len(paths))) as pool:
        file_results = list(pool.map(lambda p: _map_one_file(p, focus), paths))
    synth = synthesize_research(file_results, focus=focus)
    return {"filenames": filenames, "synth": synth, "file_results": file_results}


def research_context_block(research: dict | None) -> str:
    """Render an ingest_for_brief() result as the research section of a brief-writing prompt —
    synthesized claims (with confidence + supporting sources), not raw file text. Shared by
    brief_ai.py and brandbrief.py so the two AI-enrichment prompts read research identically."""
    if not research or not research.get("synth"):
        return ""
    synth = research["synth"]
    claims = synth.get("claims") or []
    if not claims:
        return ""
    lines = ["RESEARCH FINDINGS (synthesized across all uploaded files — cite these, do not invent "
             "beyond them; confidence reflects how many independent sources agree):"]
    for c in claims:
        supports = ", ".join(c.get("supporting_sources") or []) or "?"
        dissent = c.get("dissenting_sources") or []
        line = f"- [{c.get('confidence', '?')}] {c.get('claim', '')} (sources: {supports})"
        if dissent:
            line += f" — CONTESTED by: {', '.join(dissent)}"
        lines.append(line)
    note = synth.get("pattern_note", "")
    if note:
        lines.append(f"\nCross-source read: {note}")
    not_used = synth.get("considered_not_used") or []
    if not_used:
        lines.append("\n(Reviewed but not used — do not cite: " +
                     "; ".join(f"{c.get('source')} ({c.get('reason')})" for c in not_used) + ")")
    return "\n".join(lines)

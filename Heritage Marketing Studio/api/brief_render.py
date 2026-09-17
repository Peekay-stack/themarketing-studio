"""brief_render.py — render the full IMC brief to .docx in PURE PYTHON.

Produces the same rich document the comprehensive-brand-brief skill does — executive
summary, competitor snapshot, messaging matrix, an embedded NeedScope wheel, an embedded
CB/CA -> DB/DA diagram, the SMP callout, snapshots, opportunities/threats, recommendations
and caveats — WITHOUT Node.js or cairosvg.

Figures are drawn with Pillow (pip-install, no native cairo); the document is assembled with
python-docx. The narrative synthesis (executive summary, snapshots, recommendations, and the
optional AI enrichment) is reused from brandbrief.py so behaviour matches the skill.
"""
from __future__ import annotations

import math
import os
import tempfile

import brandbrief  # to_skill_brief(), enrich_with_ai(), _apply_enrichment()

# NeedScope wheel geometry (matches the front end's 800x640 canvas)
_CANVAS_W, _CANVAS_H = 800, 640
_CX, _CY = 400, 320
_TERR_ANGLE = {"Power": 270, "Freedom": 330, "Vitality": 30,
               "Belonging": 90, "Security": 150, "Discernment": 210}
_TERR_FILL = {"Power": "#E8736A", "Freedom": "#E89A4E", "Vitality": "#E8C547",
              "Belonging": "#5BA86B", "Security": "#4F8AC9", "Discernment": "#9A6BB8"}
_TERR_LABEL = {"Power": "#C04A40", "Freedom": "#C97829", "Vitality": "#A78318",
               "Belonging": "#3F814C", "Security": "#3469A8", "Discernment": "#7848A0"}
_GREEN = "3F814C"
_INK = "14331F"
S = 2  # render scale for crisp images


# ---------------------------------------------------------------- fonts / colour
def _font(size, bold=False):
    """`size` is in LOGICAL units, like every other measurement in this module.

    Coordinates and line spacing are all multiplied by S when drawn, so the font has to be as
    well. Without it the type was rendered at half scale against a 2x canvas, which is why the
    figures came out unreadable once shrunk to page width.
    """
    from PIL import ImageFont
    px = max(1, int(round(size * S)))
    cands = ([r"C:\Windows\Fonts\arialbd.ttf", "arialbd.ttf", "DejaVuSans-Bold.ttf"]
             if bold else
             [r"C:\Windows\Fonts\arial.ttf", "arial.ttf", "DejaVuSans.ttf"])
    for c in cands:
        try:
            return ImageFont.truetype(c, px)
        except Exception:
            continue
    return ImageFont.load_default()


def _hex(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def _tint(hexcol, toward_white=0.55):
    r, g, b = _hex(hexcol)
    return (int(r + (255 - r) * toward_white),
            int(g + (255 - g) * toward_white),
            int(b + (255 - b) * toward_white))


def _polar(angle_deg, r):
    t = math.radians(angle_deg)
    return _CX + r * math.cos(t), _CY + r * math.sin(t)


def _wrap(draw, text, font, max_w):
    words, lines, cur = str(text).split(), [], ""
    for w in words:
        trial = (cur + " " + w).strip()
        if draw.textlength(trial, font=font) <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


# ---------------------------------------------------------------- NeedScope wheel
def draw_needscope(payload: dict, out_png: str):
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (_CANVAS_W * S, _CANVAS_H * S), "white")
    d = ImageDraw.Draw(img)
    ns = payload.get("needscope", {}) or {}
    R = 220
    box = [(_CX - R) * S, (_CY - R) * S, (_CX + R) * S, (_CY + R) * S]

    # six sectors, each 60 degrees centred on its territory angle
    for terr, ang in _TERR_ANGLE.items():
        d.pieslice(box, ang - 30, ang + 30, fill=_tint(_TERR_FILL[terr]))
    d.ellipse(box, outline=(150, 150, 150), width=2)

    # territory labels around the rim
    for terr, ang in _TERR_ANGLE.items():
        lx, ly = _polar(ang, R + 34)
        f = _font(20, bold=True)
        txt = terr.upper()
        w = d.textlength(txt, font=f)
        d.text((lx * S - w / 2, ly * S - 12), txt, fill=_hex(_TERR_LABEL[terr]), font=f)

    pins = ns.get("pins", []) or []
    anchor = next((p for p in pins if p.get("anchor")), None)

    # strategic bridge (dashed line + arrow + label)
    bt = ns.get("bridge_territory")
    if bt in _TERR_ANGLE and anchor:
        tx, ty = _polar(_TERR_ANGLE[bt], 120)
        ax, ay = float(anchor["x"]), float(anchor["y"])
        _dashed(d, ax * S, ay * S, tx * S, ty * S, _hex(_GREEN), width=3, dash=14, gap=10)
        _arrowhead(d, ax * S, ay * S, tx * S, ty * S, _hex(_GREEN), size=16)
        label = str(ns.get("bridge_label", "")).strip()
        if label:
            f = _font(19, bold=True)
            mx, my = (ax + tx) / 2, (ay + ty) / 2
            w = d.textlength('"' + label + '"', font=f)
            d.text((mx * S - w / 2, my * S + 10), '"' + label + '"', fill=_hex(_GREEN), font=f)

    # Centre badge FIRST. Drawn last it painted over any pin (and its label) sitting near the
    # middle — the anchor brand is usually exactly there, so its name was being swallowed.
    br = 42 * S
    d.ellipse([_CX * S - br, _CY * S - br, _CX * S + br, _CY * S + br],
              fill="white", outline=(150, 150, 150), width=2)
    _center(d, "NeedScope", _CX * S, _CY * S - 12 * S, _font(18, bold=True), (30, 30, 30))
    # Shrink the category until it fits inside the badge rather than spilling over its edge.
    cat = str(payload.get("category", ""))
    if cat:
        f_cat = _font(14)
        for size in (14, 12, 11, 10, 9):
            f_cat = _font(size)
            if d.textlength(cat, font=f_cat) <= (2 * br - 16 * S):
                break
        _center(d, cat, _CX * S, _CY * S + 4 * S, f_cat, (120, 120, 120))

    # pins, on top of the badge
    for p in pins:
        x, y = float(p.get("x", _CX)) * S, float(p.get("y", _CY)) * S
        is_a = bool(p.get("anchor"))
        rr = (11 if is_a else 7) * S
        col = _hex("#14331F") if is_a else (119, 119, 119)
        d.ellipse([x - rr, y - rr, x + rr, y + rr], fill=col, outline="white", width=2 * S)
        f = _font(18, bold=is_a)
        name = str(p.get("brand", "")) + (" (anchor)" if is_a else "")
        w = d.textlength(name, font=f)
        dy = -34 * S / 2 if y < _CY * S else 22 * S / 2
        # A pin over the badge would still have its label lost against it — clear the badge.
        if math.hypot(x - _CX * S, y - _CY * S) < br + 26 * S:
            dy = -(br + 30 * S) if y <= _CY * S else (br + 14 * S)
        d.text((x - w / 2, y + dy), name, fill=(34, 34, 34), font=f)

    img.save(out_png)


def _dashed(d, x1, y1, x2, y2, col, width=2, dash=12, gap=8):
    dist = math.hypot(x2 - x1, y2 - y1)
    if dist == 0:
        return
    ux, uy = (x2 - x1) / dist, (y2 - y1) / dist
    pos = 0.0
    while pos < dist:
        a = pos
        b = min(pos + dash, dist)
        d.line([x1 + ux * a, y1 + uy * a, x1 + ux * b, y1 + uy * b], fill=col, width=width)
        pos += dash + gap


def _arrowhead(d, x1, y1, x2, y2, col, size=14):
    ang = math.atan2(y2 - y1, x2 - x1)
    for da in (math.radians(150), math.radians(-150)):
        d.line([x2, y2, x2 + size * math.cos(ang + da), y2 + size * math.sin(ang + da)],
               fill=col, width=3)


def _center(d, text, cx, cy, font, col):
    w = d.textlength(text, font=font)
    d.text((cx - w / 2, cy), text, fill=col, font=font)


# ---------------------------------------------------------------- CB/CA -> DB/DA
def draw_cbca(payload: dict, out_png: str):
    from PIL import Image, ImageDraw
    c = payload.get("cb_ca_db_da", {}) or {}

    # Type is sized so the figure reads at roughly 10pt once embedded at 6.3in wide. The canvas
    # height is then computed from the wrapped text rather than fixed, because larger type wraps
    # to more lines and a hard-coded card height would silently clip the longest column.
    f_title, f_label, f_body, f_decl = (_font(30, bold=True), _font(20, bold=True),
                                        _font(22), _font(21, bold=True))
    step_body, step_decl = 31, 29
    card_w, gap = 430, 120
    W = 1040
    lx, rx, top = 20, 20 + card_w + gap, 20
    wrap_w = (card_w - 40) * S

    scratch = ImageDraw.Draw(Image.new("RGB", (8, 8), "white"))

    def measure(behaviour, quote, decls) -> int:
        """Height this card's content needs, in logical units."""
        h = 52 + 26                                             # title + BEHAVIOUR label
        for line in behaviour or []:
            h += len(_wrap(scratch, "• " + str(line), f_body, wrap_w)) * step_body + 4
        h += 34                                                 # divider + ATTITUDE label
        h += len(_wrap(scratch, '"' + str(quote) + '"', f_body, wrap_w)) * step_body + 6
        for decl in (decls or []):
            if decl:
                h += len(_wrap(scratch, "— " + str(decl), f_decl, wrap_w)) * step_decl
        return h + 26                                           # bottom padding

    left = (c.get("cb", []) or [], c.get("ca_quote", "") or "", c.get("ca_declarative", []) or [])
    right = (c.get("db", []) or [], c.get("da_quote", "") or "", c.get("da_declarative", []) or [])
    card_h = max(320, measure(*left), measure(*right))
    # The from -> to caption goes in a band UNDER both cards. The gap between them is only ~120px
    # wide, so anything set there wrapped into the columns and overlapped the body text.
    shift_text = f"{c.get('shift_from','')} → {c.get('shift_to','')}".strip()
    f_shift = _font(19, bold=True)
    shift_lines = _wrap(scratch, shift_text, f_shift, (W - 120) * S) if shift_text != "→" else []
    band = (len(shift_lines) * 28 + 24) if shift_lines else 0
    H = top * 2 + card_h + band

    img = Image.new("RGB", (W * S, H * S), "white")
    d = ImageDraw.Draw(img)

    def card(x, title, tint, border, behaviour, att_quote, att_decl):
        d.rounded_rectangle([x * S, top * S, (x + card_w) * S, (top + card_h) * S],
                            radius=14 * S, fill=tint, outline=border, width=2 * S)
        _center(d, title, (x + card_w / 2) * S, (top + 12) * S, f_title, border)
        # Behaviour box
        by = top + 52
        d.text(((x + 18) * S, by * S), "BEHAVIOUR", fill=(90, 90, 90), font=f_label)
        yy = by + 26
        for line in behaviour:
            for wl in _wrap(d, "• " + str(line), f_body, wrap_w):
                d.text(((x + 20) * S, yy * S), wl, fill=(40, 50, 46), font=f_body)
                yy += step_body
            yy += 4
        # Attitude box — flows directly under the behaviour text instead of sitting at a fixed
        # offset from the bottom, so a long column can't collide with it.
        ay = yy + 8
        d.line([(x + 14) * S, ay * S, (x + card_w - 14) * S, ay * S], fill=(200, 200, 190), width=1)
        d.text(((x + 18) * S, (ay + 8) * S), "ATTITUDE", fill=(90, 90, 90), font=f_label)
        yy = ay + 34
        for wl in _wrap(d, '"' + str(att_quote) + '"', f_body, wrap_w):
            d.text(((x + 20) * S, yy * S), wl, fill=(60, 60, 60), font=f_body)
            yy += step_body
        yy += 6
        for decl in att_decl:
            if decl:
                # Wrapped, not drawn as one line — these run long and were being cut off at the
                # card edge.
                for wl in _wrap(d, "— " + str(decl), f_decl, wrap_w):
                    d.text(((x + 20) * S, yy * S), wl, fill=border, font=f_decl)
                    yy += step_decl

    card(lx, "CURRENT", _tint("#EEEEEE", 0.2), (140, 140, 140),
         c.get("cb", []) or [], c.get("ca_quote", "") or "", c.get("ca_declarative", []) or [])
    card(rx, "DESIRED", _tint("#3F814C", 0.82), _hex("#3F814C"),
         c.get("db", []) or [], c.get("da_quote", "") or "", c.get("da_declarative", []) or [])

    # centre shift arrow
    midx = (lx + card_w + gap / 2)
    ay = top + card_h / 2
    d.line([(lx + card_w + 10) * S, ay * S, (rx - 10) * S, ay * S], fill=_hex(_GREEN), width=4)
    _arrowhead(d, (lx + card_w + 10) * S, ay * S, (rx - 10) * S, ay * S, _hex(_GREEN), size=16)
    _center(d, "THE SHIFT", midx * S, (ay - 46) * S, _font(22, bold=True), _hex(_GREEN))
    for i, wl in enumerate(shift_lines):
        _center(d, wl, (W / 2) * S, (top + card_h + 14 + i * 28) * S, f_shift, _hex(_GREEN))

    img.save(out_png)


# ---------------------------------------------------------------- docx assembly
def _heading(doc, text, level, color):
    from docx.shared import RGBColor
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor.from_string(color)
    return h


def _para_hr_borders(p, color=_GREEN):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for edge in ("top", "bottom"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8")
        e.set(qn("w:space"), "8"); e.set(qn("w:color"), color)
        pbdr.append(e)
    pPr.append(pbdr)


def _table(doc, rows):
    from docx.shared import Pt
    t = doc.add_table(rows=1, cols=len(rows[0]))
    try:
        t.style = "Table Grid"
    except Exception:
        pass
    for j, cell in enumerate(rows[0]):
        c = t.rows[0].cells[j]
        c.text = str(cell)
        for para in c.paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.size = Pt(9)
    for row in rows[1:]:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = str(val)
            for para in cells[j].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
    return t


def _picture(doc, path, width_in):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches
    if not (path and os.path.exists(path)):
        return
    doc.add_picture(path, width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def build_docx(brief: dict, out_path: str):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(11)

    # Title block
    t = doc.add_paragraph()
    r = t.add_run(f"{brief.get('brand','Brand')} — IMC Brief"); r.bold = True; r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(_INK)
    sub = doc.add_paragraph()
    sr = sub.add_run(f"Prepared for: {brief.get('prepared_for','')}   ·   Date: {brief.get('date','')}")
    sr.font.size = Pt(10); sr.font.color.rgb = RGBColor.from_string("666666")
    src = doc.add_paragraph()
    # "Sources:" with nothing after it was printed on page 1 of every brief. A heading with no
    # content under it reads as a section somebody forgot to finish.
    _sources = str(brief.get("sources_note", "") or "").strip()
    scr = src.add_run(f"Sources: {_sources}" if _sources
                      else "Sources: none attached — this brief was written from the prompt alone.")
    scr.italic = True; scr.font.size = Pt(9); scr.font.color.rgb = RGBColor.from_string("666666")

    # 1. Executive summary
    _heading(doc, "1. Executive Summary", 1, "2A2A2A")
    for para in brief.get("executive_summary", []) or []:
        doc.add_paragraph(str(para))
    smp_p = doc.add_paragraph(); smp_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    smp_r = smp_p.add_run(brief.get("smp", "")); smp_r.bold = True; smp_r.italic = True; smp_r.font.size = Pt(12)

    # 2. Backgrounder — orients a reader in the evidence before any framework mechanics start. See
    # brief_skill/references/backgrounder-guidance.md for what each part is for and how it hands off
    # to the CB/CA section several sections later.
    _heading(doc, "2. Backgrounder", 1, "2A2A2A")
    bg = brief.get("backgrounder", {}) or {}
    _heading(doc, "Category perspective", 2, _GREEN)
    doc.add_paragraph(bg.get("category_perspective") or "NOT ASSESSED. No category-level data was supplied.")
    _heading(doc, "Current situation", 2, _GREEN)
    doc.add_paragraph(bg.get("current_situation")
                      or "NOT SUPPLIED. No market-share or household data was uploaded.")
    _heading(doc, "Consumer insights", 2, _GREEN)
    insights = bg.get("consumer_insights") or []
    if insights:
        for s in insights:
            doc.add_paragraph(str(s), style="List Bullet")
    else:
        doc.add_paragraph("NOT ASSESSED. No qualitative research was supplied.")
    _heading(doc, "Problem statement", 2, _GREEN)
    doc.add_paragraph(bg.get("problem_statement") or "NOT DEFINED.")

    # 3. Competitor snapshot
    _heading(doc, "3. Competitor Snapshot", 1, "2A2A2A")
    comp_rows = [["Competitor", "Hero brands", "Positioning in one line", "Recent strategic moves"]]
    for c in brief.get("competitors", []) or []:
        comp_rows.append([c.get("name", ""), c.get("hero_brands", ""),
                          c.get("positioning", ""), c.get("recent_moves", "")])
    _table(doc, comp_rows)

    # 4. Messaging comparison matrix
    _heading(doc, "4. Messaging Comparison Matrix", 1, "2A2A2A")
    mm = brief.get("messaging_matrix", {}) or {}
    dims, brands, cells = mm.get("dimensions", []), mm.get("brands", []), mm.get("cells", [])
    if dims and brands:
        mrows = [["Dimension"] + list(brands)]
        for i, dim in enumerate(dims):
            row = [dim] + (list(cells[i]) if i < len(cells) else [])
            mrows.append(row + [""] * (len(brands) + 1 - len(row)))
        _table(doc, mrows)

    # 5. NeedScope
    _heading(doc, "5. NeedScope Brand Positioning", 1, "2A2A2A")
    doc.add_paragraph("NeedScope maps brands across six colour-coded emotional territories arranged "
                      "on a wheel. The focal brand is anchored and a dashed bridge shows the strategic "
                      "move into an adjacent territory.")
    ns = brief.get("needscope", {}) or {}
    _picture(doc, ns.get("figure_path"), 5.5)
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run("Figure 1. NeedScope wheel and strategic bridge."); cr.italic = True
    cr.font.size = Pt(9); cr.font.color.rgb = RGBColor.from_string("666666")
    _heading(doc, "Reading the wheel", 2, _GREEN)
    doc.add_paragraph(ns.get("reading", ""))
    _heading(doc, "Strategic implication", 2, _GREEN)
    doc.add_paragraph(ns.get("strategic_implication", ""))

    # 6. CB/CA -> DB/DA
    _heading(doc, "6. Messaging Strategy — Current → Desired", 1, "2A2A2A")
    doc.add_paragraph("The messaging job framed as the shift the work must deliver.")
    cc = brief.get("cb_ca_db_da", {}) or {}
    _picture(doc, cc.get("figure_path"), 6.3)
    cap2 = doc.add_paragraph(); cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr2 = cap2.add_run("Figure 2. Current → Desired behaviour & attitude."); cr2.italic = True
    cr2.font.size = Pt(9); cr2.font.color.rgb = RGBColor.from_string("666666")
    shift_p = doc.add_paragraph()
    shift_p.add_run("The shift in one line: ").font.size = Pt(11)
    sb = shift_p.add_run(cc.get("shift_line", "")); sb.bold = True
    if cc.get("caption"):
        doc.add_paragraph(cc.get("caption"))

    # 7. SMP
    _heading(doc, "7. Single Minded Proposition", 1, "2A2A2A")
    callout = doc.add_paragraph(); callout.alignment = WD_ALIGN_PARAGRAPH.CENTER
    co = callout.add_run(brief.get("smp", "")); co.bold = True; co.italic = True
    co.font.size = Pt(16); co.font.color.rgb = RGBColor.from_string(_GREEN)
    _para_hr_borders(callout)
    _heading(doc, f"Why only {brief.get('brand','the brand')} can say it", 2, _GREEN)
    for defn in brief.get("smp_defence", []) or []:
        doc.add_paragraph(f"{defn.get('competitor','')}: {defn.get('why_collapses','')}", style="List Bullet")
    _heading(doc, "What the SMP unlocks (illustrative — not final copy)", 2, _GREEN)
    un = brief.get("smp_unlocks", {}) or {}
    for label, key in (("Brand line", "brand_line"), ("Mass line", "mass_line"),
                       ("Activation", "activation"), ("Product", "product")):
        if un.get(key):
            doc.add_paragraph(f"{label}: {un.get(key)}", style="List Bullet")

    # 8. Pricing / distribution / content
    _heading(doc, "8. Pricing, Distribution & Content Snapshot", 1, "2A2A2A")
    snap = brief.get("snapshots", {}) or {}
    for label, key in (("Pricing", "pricing"), ("Distribution", "distribution"),
                       ("Content & campaigns", "content_campaigns"),
                       ("Content gaps to claim now", "content_gaps")):
        _heading(doc, label, 2, _GREEN)
        doc.add_paragraph(snap.get(key, ""))

    # 9. Opportunities & threats
    _heading(doc, "9. Opportunities & Threats", 1, "2A2A2A")
    ot_rows = [["Opportunities", "Threats"]]
    for r_ in brief.get("opportunities_threats", []) or []:
        ot_rows.append([r_.get("opportunity", ""), r_.get("threat", "")])
    if len(ot_rows) > 1:
        _table(doc, ot_rows)

    # 10. Recommended actions
    _heading(doc, "10. Recommended Actions", 1, "2A2A2A")
    rec = brief.get("recommendations", {}) or {}
    if rec.get("preface"):
        pf = doc.add_paragraph(); pr = pf.add_run(rec.get("preface")); pr.italic = True
        pr.font.color.rgb = RGBColor.from_string("555555")
    _heading(doc, "Quick wins (this quarter)", 2, _GREEN)
    for s in rec.get("quick_wins", []) or []:
        doc.add_paragraph(s, style="List Bullet")
    _heading(doc, "Strategic moves (12–18 months)", 2, _GREEN)
    for s in rec.get("strategic_moves", []) or []:
        doc.add_paragraph(s, style="List Number")

    # 11. Caveats
    _heading(doc, "11. Caveats", 1, "2A2A2A")
    doc.add_paragraph(brief.get("caveats", ""))
    foot = doc.add_paragraph(); fr = foot.add_run("Prepared with the IMC Brief builder.")
    fr.italic = True; fr.font.size = Pt(9); fr.font.color.rgb = RGBColor.from_string("777777")

    doc.save(out_path)


def generate(payload: dict, research: dict | None = None, workdir: str | None = None) -> str:
    """Build the rich .docx and return its path — pure Python, no Node/cairosvg.

    research: research_parse.ingest_for_brief() output for any uploaded files, or None.
    """
    workdir = workdir or tempfile.mkdtemp(prefix="imcbrief_")
    ns_png = os.path.join(workdir, "needscope.png")
    cbca_png = os.path.join(workdir, "cbca.png")
    # figures are best-effort: if Pillow is somehow unavailable, still produce the text+tables
    try:
        draw_needscope(payload, ns_png)
    except Exception:
        ns_png = None
    try:
        draw_cbca(payload, cbca_png)
    except Exception:
        cbca_png = None
    brief = brandbrief.to_skill_brief(payload, ns_png or "", cbca_png or "")
    # sources_note used to come only from payload.get("sources_note", "") — nothing in this pipeline
    # ever set that key, so "Sources: none attached" printed even when files WERE attached (the
    # contradiction the user's A/B test caught, alongside "Per the uploaded market data" two lines
    # later). Set it here from the actual filenames the route was given, not a model self-report.
    filenames = (research or {}).get("filenames") or []
    if filenames:
        note = f"{', '.join(filenames)}."
        brief["sources_note"] = note
        brief["caveats"] = (note + " NeedScope positions are as placed in the builder; validate "
                            "market-share, pricing and competitor activity before externalising.")
    ai = brandbrief.enrich_with_ai(payload, research)
    if ai:
        brief = brandbrief._apply_enrichment(brief, payload, ai)
    elif os.environ.get("ANTHROPIC_API_KEY"):
        # A key IS configured, so `ai is None` here means the call actually failed (not the benign
        # "no key" case, which returns the same None — see enrich_with_ai()'s own comment on why that
        # used to be indistinguishable). Executive summary, the pricing/distribution snapshot,
        # opportunities & threats and recommended actions are all still sitting at their structural
        # defaults with nothing said about it — a real export did exactly this, silently, and it took
        # a line-by-line diff against brandbrief.py's own hardcoded strings to even notice. Say so in
        # the document itself rather than let a reader mistake placeholder text for a tailored read.
        note = ("AI enrichment could not complete for this export — the executive summary, pricing/"
                "distribution snapshot, opportunities & threats and recommended actions below are "
                "structural defaults, not a tailored read of the attached research. Redraft or "
                "regenerate to retry.")
        brief["caveats"] = (brief.get("caveats", "") + " " + note).strip()
    # ensure the figure paths point at the PNGs we actually drew
    brief.setdefault("needscope", {})["figure_path"] = ns_png
    brief.setdefault("cb_ca_db_da", {})["figure_path"] = cbca_png
    base = "".join(ch if ch.isalnum() else "_" for ch in str(payload.get("brand", "brand"))).strip("_") or "brand"
    out_docx = os.path.join(workdir, f"{base}_IMC_Brief.docx")
    build_docx(brief, out_docx)
    return out_docx

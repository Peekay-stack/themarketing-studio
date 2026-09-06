"""roundtrip.py — read an edited Word document back into the studio's own structure.

The studio writes briefs and scripts to Word so they can be worked on the way marketing teams
actually work on them: offline, in tracked changes, by people who will never open this app. The
return leg was missing, so that editing was a dead end — the portal's copy stayed stale and the
rewrite lived only in someone's Downloads folder.

Two paths back in, tried in order:

  1. **Structural.** Documents this app produced have a known shape — the script is one table with
     SCENE / VISUAL / CAMERA / DIALOGUE / SUPER columns, the brief is heading-then-body. If that
     shape survived the edit, the text is read straight out of it. Exact, free, and offline.
  2. **Interpretive.** If the shape did not survive — someone restructured it, wrote it fresh, or
     sent a document from elsewhere — the text is handed to the model with the target schema.

Which one ran is always reported, because a structural read is a fact and an interpreted one is a
best effort, and the person approving it deserves to know which they are looking at.
"""
from __future__ import annotations

import io
import os
import re

import jsonout

# The script table as this app writes it (see docs.py). Matching is loose on case and spacing so a
# reformatted header still lands.
_SCRIPT_COLS = ("scene", "visual", "camera", "dialogue", "super")


def read_docx(data: bytes) -> dict:
    """Pull everything readable out of a .docx: paragraphs, tables, and the flat text."""
    from docx import Document
    doc = Document(io.BytesIO(data))
    paras = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    tables = []
    for t in doc.tables:
        rows = []
        for r in t.rows:
            rows.append([c.text.strip() for c in r.cells])
        if rows:
            tables.append(rows)
    flat = "\n".join(paras)
    for tb in tables:
        flat += "\n" + "\n".join(" | ".join(r) for r in tb)
    return {"paragraphs": paras, "tables": tables, "text": flat.strip()}


# --- scripts ---------------------------------------------------------------------------------

def _header_map(row: list[str]) -> dict[str, int]:
    """Column name -> index, for a header row that may have been renamed slightly."""
    found = {}
    for i, cell in enumerate(row):
        c = re.sub(r"[^a-z]", "", (cell or "").lower())
        for key in _SCRIPT_COLS:
            if key in c and key not in found:
                found[key] = i
    return found


def script_from_tables(tables: list[list[list[str]]]) -> list[dict]:
    """Rows of the studio's script table, or [] if no table looks like one."""
    for tb in tables:
        if len(tb) < 2:
            continue
        cols = _header_map(tb[0])
        if "visual" not in cols or "dialogue" not in cols:
            continue                       # not the script table
        out = []
        for n, row in enumerate(tb[1:], start=1):
            def cell(key: str) -> str:
                i = cols.get(key, -1)
                return row[i].strip() if 0 <= i < len(row) else ""
            visual, dialogue = cell("visual"), cell("dialogue")
            if not visual and not dialogue:
                continue                   # spacer row
            # "3" or "3\n8–12s" — the number and the timecode share a cell in our export.
            head = cell("scene")
            num = re.search(r"\d+", head)
            tc = re.search(r"\d+\s*[-–—]\s*\d+\s*s", head)
            out.append({
                "no": int(num.group()) if num else n,
                "tc": tc.group().replace(" ", "") if tc else "",
                "visual": visual, "camera": cell("camera"),
                "dialogue": dialogue, "supr": cell("super"),
            })
        if out:
            return out
    return []


def script_from_text(text: str) -> list[dict]:
    """Last resort for a script that is prose: ask the model to lay it back out."""
    schema = ("A JSON array. One object per scene, in order, with exactly these keys: "
              '"no" (integer), "visual" (what the camera sees), "camera" (the shot), '
              '"dialogue" (exactly what is said, including who says it, e.g. "Mother: Ghar jaisa." '
              'or "VO: ..."), "supr" (on-screen text, "" if none).')
    rows = _ask(text, schema, "film script")
    out = []
    for i, r in enumerate(rows or [], start=1):
        if not isinstance(r, dict):
            continue
        out.append({"no": int(r.get("no") or i), "tc": "",
                    "visual": str(r.get("visual") or ""), "camera": str(r.get("camera") or ""),
                    "dialogue": str(r.get("dialogue") or ""), "supr": str(r.get("supr") or "")})
    return out


def import_script(data: bytes) -> dict:
    """An edited script document -> rows the Shoot Board and Production screens can use."""
    doc = read_docx(data)
    rows = script_from_tables(doc["tables"])
    how = "structure"
    if not rows:
        rows = script_from_text(doc["text"])
        how = "interpreted"
    notes = []
    if not rows:
        return {"ok": False, "rows": [], "how": how,
                "detail": "Couldn't find a script in that document — no scene table and no scenes "
                          "the model could lay out. Is it the right file?"}
    # Timecodes are deliberately dropped: scene lengths are re-derived from the dialogue by
    # filmvoice.retime(), so importing stale ones would fight the timing the render actually uses.
    for r in rows:
        r["tc"] = ""
    notes.append(f"{len(rows)} scenes read "
                 + ("straight from the script table" if how == "structure"
                    else "by interpreting the document — check them before approving"))
    notes.append("Scene lengths will be re-timed from the dialogue when you open the Shoot Board")
    return {"ok": True, "rows": rows, "how": how, "detail": " · ".join(notes)}


# --- briefs ----------------------------------------------------------------------------------

# The editable fields on the IMC screen. Kept here so the import can only ever fill fields that
# actually exist — an invented key would silently vanish at render time.
BRIEF_FIELDS = [
    ("businessObjective", "The business result the work must deliver"),
    ("marketingObjective", "The marketing objective"),
    ("background", "Background and context"),
    ("targetAudience", "Who this speaks to"),
    ("consumerInsight", "The consumer insight"),
    ("currentBelief", "Current belief / current attitude (CB / CA)"),
    ("desiredBelief", "Desired belief / desired attitude (DB / DA)"),
    ("smp", "The single-minded proposition"),
    ("rtbs", "Reasons to believe"),
    ("toneOfVoice", "Tone of voice"),
    ("mandatories", "Mandatories"),
    ("packHierarchy", "Pack / communication hierarchy"),
    ("successMetrics", "How success is measured"),
]


def brief_from_headings(doc: dict) -> dict:
    """Match the document's headings against the brief's fields.

    Our own export writes each field under its label, so an edit that kept the labels round-trips
    exactly — including any text the user added underneath.
    """
    paras = doc["paragraphs"]
    labels = {}
    for key, label in BRIEF_FIELDS:
        labels[re.sub(r"[^a-z]", "", label.lower())] = key
        labels[re.sub(r"[^a-z]", "", key.lower())] = key
    out, current, buf = {}, None, []

    def flush():
        if current and buf:
            out[current] = "\n".join(buf).strip()

    for p in paras:
        norm = re.sub(r"[^a-z]", "", p.lower())
        # A heading is short and matches a field name; anything longer is body copy.
        hit = labels.get(norm) if len(p) < 80 else None
        if hit:
            flush()
            current, buf = hit, []
            continue
        if current:
            buf.append(p)
    flush()
    return out


def import_brief(data: bytes) -> dict:
    """An edited brief document -> the IMC screen's fields."""
    doc = read_docx(data)
    fields = brief_from_headings(doc)
    how = "structure"
    if len(fields) < 3:                    # too little matched to call it a structural read
        schema = ("A JSON object. Use only these keys, omitting any the document does not cover: "
                  + ", ".join(f'"{k}"' for k, _ in BRIEF_FIELDS)
                  + ". Each value is the text from the document for that section, verbatim where "
                    "possible — do not summarise, invent or improve it.")
        got = _ask(doc["text"], schema, "marketing brief")
        if isinstance(got, dict):
            allowed = {k for k, _ in BRIEF_FIELDS}
            fields = {k: str(v).strip() for k, v in got.items() if k in allowed and str(v).strip()}
            how = "interpreted"
    if not fields:
        return {"ok": False, "fields": {}, "how": how,
                "detail": "Couldn't find brief sections in that document."}
    named = ", ".join(sorted(fields))
    return {"ok": True, "fields": fields, "how": how,
            "detail": (f"{len(fields)} sections read "
                       + ("from the document's own headings" if how == "structure"
                          else "by interpreting the document — read them before approving")
                       + f": {named}")}


# --- the interpretive path ---------------------------------------------------------------------

def _ask(text: str, schema: str, what: str):
    """Ask the model to fit free text to a schema. Returns parsed JSON, or None."""
    if not os.environ.get("ANTHROPIC_API_KEY") or not text.strip():
        return None
    import anthropic
    prompt = (
        f"Below is a {what} that someone rewrote outside this system. Read it and return its "
        f"content in the required shape.\n\n"
        f"REQUIRED OUTPUT: {schema}\n\n"
        "Rules: return JSON and nothing else — no preamble, no code fence. Use the document's own "
        "words; do not summarise, improve, or add anything that is not there. If the document is "
        "silent on something, leave it out rather than inventing it.\n\n"
        f"--- DOCUMENT ---\n{text[:24000]}"
    )
    try:
        client = anthropic.Anthropic()
        msg = client.messages.create(
            model=os.environ.get("GEN_MODEL", "claude-opus-4-8"),
            max_tokens=8000, messages=[{"role": "user", "content": prompt}])
        raw = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text").strip()
    except Exception as e:                 # a failed import must not take the request down
        print(f"[roundtrip] model read failed: {type(e).__name__}: {e}")
        return None
    raw = re.sub(r"^```(?:json)?|```$", "", raw.strip(), flags=re.MULTILINE).strip()
    start = min([i for i in (raw.find("{"), raw.find("[")) if i >= 0], default=-1)
    if start > 0:
        raw = raw[start:]
    try:
        data, _err = jsonout.extract(raw)
        return data if data is not None else {}
    except ValueError:
        print("[roundtrip] model returned something that was not JSON")
        return None

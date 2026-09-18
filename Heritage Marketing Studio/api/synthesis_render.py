"""synthesis_render.py — turns synthesis.build_linkages()'s output into the standalone .pptx.

Deliberately plain (see synthesis_skill/references/deck-structure.md's "Visual treatment" section): this
deck earns trust through what it says, not production values. Same green/ink palette as brief_render.py's
.docx for visual family resemblance across the product's document outputs, no charts fabricated from
anything that wasn't already computed by research_parse.py's real spreadsheet aggregation.
"""
from __future__ import annotations

import os
import time

_GREEN = "3F814C"
_INK = "14331F"
_GREY = "666666"
_LIGHT = "F2F5F2"

_CONF_LABEL = {"high": "High confidence", "medium": "Medium confidence", "low": "Low confidence", "": ""}


def _rgb(hexcol: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexcol)


_FOOTER_MARK = os.path.join(os.path.dirname(__file__), "frontend", "assets", "tms-footer-mark.png")


def _add_slide_footer(slide):
    """The Marketing Studio's own mark, bottom-left of every slide — this deck leaves the portal and
    gets circulated on its own. The mark (tms-footer-mark.png) already renders "the marketing studio"
    as part of the lockup, so no separate URL caption runs next to it — the user asked for the logo
    alone, a touch larger, after the first version doubled up and crowded the two together. Same asset
    and treatment as the docx footers (brief_render.py/docs.py's `_add_footer`)."""
    from pptx.util import Inches
    if os.path.exists(_FOOTER_MARK):
        try:
            slide.shapes.add_picture(_FOOTER_MARK, Inches(0.4), Inches(6.98), height=Inches(0.28))
        except Exception:
            pass


def _title_slide(prs, title: str, subtitle: str):
    from pptx.util import Pt
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(32)
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = _rgb(_INK)
    if len(slide.placeholders) > 1:
        sub = slide.placeholders[1]
        sub.text = subtitle
        for p in sub.text_frame.paragraphs:
            p.font.size = Pt(13)
            p.font.color.rgb = _rgb(_GREY)
    _add_slide_footer(slide)


def _new_slide(prs, title: str, title_color: str = _INK):
    from pptx.util import Inches, Pt
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = title
    tf = slide.shapes.title.text_frame
    tf.paragraphs[0].font.size = Pt(24)
    tf.paragraphs[0].font.color.rgb = _rgb(title_color)
    tf.word_wrap = True
    _add_slide_footer(slide)
    return slide


def _body_box(slide, top_in: float = 1.6, height_in: float = 5.2):
    from pptx.util import Inches
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top_in), Inches(9.0), Inches(height_in))
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def _bullet(tf, text: str, *, size: int = 15, bold: bool = False, color: str = _INK,
           first: bool = False, level: int = 0, space_after: int = 10):
    from pptx.util import Pt
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.level = level
    p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = _rgb(color)
    return p


def _confidence_chip(tf, confidence: str, *, first: bool = False):
    label = _CONF_LABEL.get(str(confidence or "").lower(), "")
    if not label:
        return
    _bullet(tf, label, size=12, bold=True, color=_GREEN, first=first, space_after=6)


def _headline_slide(prs, headline: dict, standalone: list[dict]):
    slide = _new_slide(prs, "The read")
    tf = _body_box(slide)
    has_pattern = bool(headline.get("has_pattern"))
    _confidence_chip(tf, headline.get("confidence", ""), first=True)
    _bullet(tf, headline.get("statement", ""), size=17, bold=True, color=_INK,
           first=not tf.paragraphs[0].runs, space_after=16)
    if not has_pattern and standalone:
        _bullet(tf, "What each source says on its own:", size=13, bold=True, color=_GREY, space_after=8)
        for s in standalone[:8]:
            srcs = ", ".join(s.get("sources") or []) or "?"
            _bullet(tf, f"{s.get('text', '')}  —  {srcs}", size=13, level=1, color=_INK, space_after=6)


def _linkage_slide(prs, linkage: dict, index: int):
    label = str(linkage.get("label") or "").strip()
    title = linkage.get("title", "") or f"Linkage {index}"
    slide = _new_slide(prs, f"{label + ': ' if label else ''}{title}")
    tf = _body_box(slide, top_in=1.7, height_in=4.3)
    claims = linkage.get("claims") or []
    for i, c in enumerate(claims):
        srcs = ", ".join(c.get("sources") or []) or "?"
        conf = _CONF_LABEL.get(str(c.get("confidence", "")).lower(), "")
        prefix = "↓ " if i > 0 else ""
        _bullet(tf, f"{prefix}{c.get('text', '')}", size=14, first=(i == 0), space_after=4)
        _bullet(tf, f"     {conf} · {srcs}" if conf else f"     {srcs}", size=11, color=_GREY, space_after=12)
    # The implication, set apart in a bordered box — the one thing on the slide that isn't a claim.
    from pptx.util import Inches, Pt
    box_top = 1.7 + 4.3 + 0.15
    impl = slide.shapes.add_textbox(Inches(0.6), Inches(box_top), Inches(9.0), Inches(0.9))
    impl.fill.solid(); impl.fill.fore_color.rgb = _rgb(_LIGHT)
    impl.line.color.rgb = _rgb(_GREEN); impl.line.width = Pt(1)
    tfi = impl.text_frame; tfi.word_wrap = True
    tfi.margin_left = Inches(0.15); tfi.margin_top = Inches(0.08)
    p = tfi.paragraphs[0]
    r = p.add_run(); r.text = f"Points to: {title}"
    r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = _rgb(_GREEN)
    conf = _CONF_LABEL.get(str(linkage.get("confidence", "")).lower(), "")
    if conf:
        p2 = tfi.add_paragraph()
        r2 = p2.add_run(); r2.text = conf
        r2.font.size = Pt(11); r2.font.color.rgb = _rgb(_GREY)


def _considered_slide(prs, considered: list[dict]):
    slide = _new_slide(prs, "Considered, not used")
    tf = _body_box(slide)
    for i, c in enumerate(considered[:10]):
        _bullet(tf, f"{c.get('source', '')}", size=14, bold=True, first=(i == 0), space_after=2)
        _bullet(tf, f"     {c.get('reason', '')}", size=12, color=_GREY, space_after=10)


def _caveats_slide(prs, filenames: list[str], pattern_note: str):
    slide = _new_slide(prs, "Caveats & sourcing")
    tf = _body_box(slide)
    _bullet(tf, "Sources reviewed:", size=13, bold=True, first=True, space_after=4)
    _bullet(tf, ", ".join(filenames) or "(none)", size=12, color=_GREY, level=1, space_after=14)
    if pattern_note:
        _bullet(tf, "Cross-source read, in full:", size=13, bold=True, space_after=4)
        _bullet(tf, pattern_note, size=12, color=_GREY, level=1, space_after=14)
    _bullet(tf, "Validate every figure above against the latest live data before external use — this "
               "deck reads whatever was uploaded, verbatim, and carries forward any source's own "
               "disclosure that its numbers are illustrative or simulated.", size=11, color=_GREY)


def generate(deck: dict, out_path: str | None = None, workdir: str | None = None) -> str:
    """Build the .pptx and return its path.

    deck: synthesis.build_linkages() output — {filenames, headline, linkages, standalone_findings,
    considered_not_used, pattern_note}.
    """
    from pptx import Presentation
    prs = Presentation()

    filenames = deck.get("filenames") or []
    title = deck.get("title") or "Cross-Source Research Synthesis"
    subtitle = (", ".join(filenames) or "No sources listed") + f"\n{time.strftime('%d %b %Y')}"
    _title_slide(prs, title, subtitle)

    _headline_slide(prs, deck.get("headline") or {}, deck.get("standalone_findings") or [])

    for i, link in enumerate(deck.get("linkages") or [], 1):
        _linkage_slide(prs, link, i)

    considered = deck.get("considered_not_used") or []
    if considered:
        _considered_slide(prs, considered)

    _caveats_slide(prs, filenames, deck.get("pattern_note", ""))

    workdir = workdir or __import__("tempfile").mkdtemp(prefix="synthesis_")
    out_path = out_path or os.path.join(workdir, "Research_Synthesis.pptx")
    prs.save(out_path)
    return out_path
